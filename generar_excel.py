"""
Generador de base de datos de categorías MercadoLibre Argentina — Autopartes
Fuente de datos: capturas de pantalla del árbol de categorías de MeLi AR
Uso: python generar_excel.py

Convención de filas: (N1, N2, N3, N4, N5, URL, ID_MLA, Notas)
  - URL: URL verificada desde captura (✓) o inferida por slug (*)
  - Dejar en "" los niveles/campos que no apliquen
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_FILE = "categorias_autopartes_meli.xlsx"

HEADERS = [
    "Categoria_Nivel_1",
    "Categoria_Nivel_2",
    "Categoria_Nivel_3",
    "Categoria_Nivel_4",
    "Categoria_Nivel_5",
    "URL_Categoria",       # ✓ verificada | * inferida por slug
    "ID_Categoria_MLA",
    "Notas",
]

# ─────────────────────────────────────────────
# Helpers de URL
# Patrón verificado: listado.mercadolibre.com.ar/{n1-slug}/{n2-slug}/{n3-slug}/
# ─────────────────────────────────────────────
_BASE = "https://listado.mercadolibre.com.ar/accesorios-vehiculos/repuestos-autos-camionetas"

def url_n3(n3_slug):
    return f"{_BASE}/{n3_slug}/"

def url_n4(n3_slug, n4_slug):
    return f"{_BASE}/{n3_slug}/{n4_slug}/"


# ─────────────────────────────────────────────
# DATOS
# ─────────────────────────────────────────────
N1 = "Accesorios para Vehículos"
N2 = "Repuestos Autos y Camionetas"

CATEGORIAS = [

    # ── Baterías — terminal (sin N4) — URL inferida por slug ─────────────────
    (N1, N2, "Baterías", "", "", url_n3("baterias"), "", "* URL inferida | Terminal — va directo a productos"),

    # ── Carrocería — 37 subcategorías N4 — URL N3 verificada ✓ ───────────────
    # URL verificada: https://listado.mercadolibre.com.ar/accesorios-vehiculos/repuestos-autos-camionetas/carroceria/
    (N1, N2, "Carrocería", "Alojamientos de Ópticas",    "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Barras",                      "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Baúles",                      "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Baúles y Tapas Traseras",     "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Bisagras",                    "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Capot",                       "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Cerraduras",                  "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Componentes Adicionales",     "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Espejos Retrovisores y Vidrios", "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Grampas",                     "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Guardabarros",                "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Laterales de Carrocería",     "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Manijas",                     "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Molduras",                    "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Otros",                       "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Paneles",                     "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Paneles de Cuarto",           "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Paneles del Parabrisas",      "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Paragolpes",                  "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Parrillas",                   "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Pasaruedas",                  "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Pernos",                      "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Pilares",                     "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Pisos",                       "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Puertas y Paneles",           "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Rejillas para Faros Auxiliares", "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Reparación de Cajas",         "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Soportes Elevación del Capó", "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Soportes de Rueda de Auxilio","", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Soportes y Abrazaderas",      "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Tanques de Combustible",      "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Techos",                      "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Techos Corredizos y Motores", "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Tomas de Aire",               "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Varillas de Apoyo",           "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Ventilaciones",               "", url_n3("carroceria"), "", ""),
    (N1, N2, "Carrocería", "Zócalos",                     "", url_n3("carroceria"), "", ""),

    # ── Pendientes de captura ─────────────────────────────────────────────────
    (N1, N2, "Cerraduras y Llaves",          "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Climatización",                "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Componentes de Seguridad",     "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Conducción Asistida Avanzada", "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Electroventiladores",          "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Eléctricos, Híbridos y PHEV",  "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Encendido",                    "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Escapes",                      "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Filtros",                      "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Frenos",                       "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Iluminación",                  "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Instalaciones Eléctricas",     "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Inyección",                    "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Motor",                        "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Otros",                        "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Repuestos de Exterior",        "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Repuestos de Habitáculo",      "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Suspensión y Dirección",       "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Transmisión",                  "", "", "", "", "Pendiente captura N4"),
    (N1, N2, "Ventanas y Sellos",            "", "", "", "", "Pendiente captura N4"),
]


# ─────────────────────────────────────────────
# Estilos
# ─────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1A56DB")
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

LEVEL_COLORS = {
    0: "DBEAFE",  # N1
    1: "EFF6FF",  # N2
    2: "F8FAFF",  # N3
    3: "FFFFFF",  # N4
    4: "FFFFFF",  # N5
    5: "E0F2FE",  # URL  — celeste suave
    6: "FEF9C3",  # ID   — amarillo
    7: "F3F4F6",  # Notas — gris
}

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

COL_WIDTHS = [32, 32, 32, 30, 25, 70, 18, 45]


def build_excel(categorias, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Categorias MeLi AR"

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border    = THIN_BORDER

    ws.row_dimensions[1].height = 30

    for row_idx, row_data in enumerate(categorias, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill      = PatternFill("solid", fgColor=LEVEL_COLORS[col_idx - 1])
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border    = THIN_BORDER
        ws.row_dimensions[row_idx].height = 20

    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    ws.freeze_panes = "A2"

    wb.save(output_file)
    total = len(categorias)
    pendientes = sum(1 for r in categorias if "Pendiente" in r[-1])
    print(f"[OK] {output_file}  |  {total} filas  |  {total - pendientes} completas  |  {pendientes} pendientes")


if __name__ == "__main__":
    build_excel(CATEGORIAS, OUTPUT_FILE)
